"""Build the ten-page FlowSight project report.

Run build_metrics.py first so that the report uses measured project metrics.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "report_assets"
OUTPUT = ROOT / "report_output" / "FlowSight_Traffic_Analytics_Report.docx"

BLUE = "173F73"
LIGHT_BLUE = "EAF1FB"
PALE = "F6F8FC"
GRID = "D9D9D9"
TEXT = "1F2937"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{name}")
        edge = borders.find(tag)
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def write_para(doc, text="", bold_lead=None, size=10.5, alignment=None, before=0, after=5, italic=False):
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=size, bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), size=size, italic=italic)
    else:
        set_run_font(p.add_run(text), size=size, italic=italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 4)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12.5, bold=True, color="000000")
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run(text), size=8.5, italic=True, color="53657D")
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, title in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(str(title)), size=font_size, bold=True, color="FFFFFF")
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            set_run_font(p.add_run(str(value)), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(doc, filename, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    set_run_font(p.add_run("FlowSight Traffic Pattern and Congestion Analytics  |  "), size=8, color="64748B")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.62)
    section.bottom_margin = Inches(.56)
    section.left_margin = Inches(.70)
    section.right_margin = Inches(.70)
    footer(section)
    style = doc.styles["Normal"]
    style.font.name = "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor.from_string(TEXT)
    props = doc.core_properties
    props.title = "FlowSight Traffic Pattern and Congestion Analytics"
    props.subject = "Machine learning project report"
    props.author = "Student"


def main():
    metrics = list(csv.DictReader((ASSETS / "model_metrics.csv").open()))
    facts = json.loads((ASSETS / "metrics.json").read_text())
    doc = Document()
    setup(doc)

    # PAGE 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(90)
    set_run_font(p.add_run("FlowSight Traffic Pattern and Congestion Analytics"), size=25, bold=True, color="000000")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(28)
    set_run_font(p.add_run("An Interactive Machine Learning Dashboard for Exploring Simulated Urban Traffic Patterns"), size=13, color="4B5E76")
    write_para(doc, "PROJECT REPORT", size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=18)
    heading(doc, "Abstract", 1)
    write_para(doc, "This report presents FlowSight, an interactive traffic-pattern analytics system designed for a classroom machine learning project. Instead of claiming to predict a future journey, the application lets a user enter a road situation and inspect comparable observations. A reproducible dataset of 2,400 simulated traffic observations is analysed using K-Means clustering, K-Nearest Neighbours, a Decision Tree and Linear Regression. The dashboard updates traffic groups, similar records, interpretable decision rules and factor comparisons whenever the user changes volume, speed, rainfall, incidents, lanes or time of day. A 75:25 held-out evaluation compares four regression baselines. Linear Regression achieved the lowest test RMSE in this controlled dataset, while K-Means produced three visually distinct traffic groups with a silhouette score of 0.451. The project demonstrates data cleaning, transparent analytics and a React plus Flask deployment pathway, while clearly separating simulated evidence from real city-sensor findings.")
    write_para(doc, "Keywords: traffic analytics; congestion patterns; K-Means; KNN; decision tree; linear regression; React dashboard; explainable machine learning.", bold_lead="Keywords: ", size=10)
    write_para(doc, "Prepared for: Machine Learning Project Submission", size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER, before=38, after=4)
    write_para(doc, "Date: October 2026", size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # PAGE 2
    heading(doc, "1 Introduction")
    write_para(doc, "Urban traffic changes with demand, weather, road capacity and incidents. A simple average can hide important variation: two observations may have similar volumes but different delay because rainfall, an incident or fewer open lanes changes how the road behaves. FlowSight is a small, explainable analytics application that helps a user explore these differences through comparable recorded situations.")
    write_para(doc, "The project is deliberately an analytics tool rather than a future-traffic prediction system. A user enters six values - time of day, vehicle volume, average speed, rainfall, incidents and open lanes. The application returns a neighbourhood of similar recorded observations and makes their patterns visible. This scope makes the demonstration easier to explain and avoids presenting an illustrative classroom dataset as a live operational forecast.")
    heading(doc, "1.1 Problem statement", 2)
    write_para(doc, "Traffic datasets often combine measurements that have different scales and meanings. Analysts need a clear way to inspect which traffic state a situation resembles, which records support that conclusion, and how selected factors differ from a reference condition. A dashboard should therefore be interactive, responsive and explainable rather than only returning a single opaque score.")
    heading(doc, "1.2 Objectives", 2)
    add_table(doc, ["Objective", "Implementation evidence"], [
        ["Prepare a reliable dataset", "Range checks, duplicate checks, numeric conversion, feature selection and standardisation."],
        ["Discover traffic patterns", "K-Means groups comparable traffic states into free flow, moderate and severe congestion."],
        ["Retrieve comparable records", "KNN retrieves 10, 30, 60 or 100 closest standardized observations."],
        ["Explain results simply", "Linear contribution bars and decision-tree rules are shown beside the data."],
        ["Support demonstration", "React inputs, live charts, record filters and a Flask API run locally."],
    ], widths=[2.05, 4.85], font_size=8.8)
    heading(doc, "1.3 Scope and ethical note", 2)
    write_para(doc, "The 2,400 rows are synthetic but reproducible observations generated with seed 42. They mimic plausible relationships among volume, speed, rainfall, incidents, road lanes and delay. Results therefore validate the software workflow and the relative behaviour of algorithms in this dataset; they do not establish performance for a particular city. A production system would require authorised sensor data, data-governance controls, regular retraining and external validation.")
    page_break(doc)

    # PAGE 3
    heading(doc, "2 Literature Review")
    write_para(doc, "Recent traffic research shows a shift from single-model forecasting toward comparative, hybrid and explainable workflows. The studies below are selected because they are recent and directly relevant to clustering, classical machine learning or practical model comparison. Their datasets and targets differ, so their reported scores are not directly transferable to FlowSight.")
    add_table(doc, ["Study", "Data and approach", "Key finding", "Relevance to FlowSight"], [
        ["Lahijanian, Isaenko and Fusco (2025) [1]", "Eleven months of Italian motorway data; K-Means, DBSCAN and affinity propagation.", "Six day types were identified; K-Means and DBSCAN were compatible while affinity propagation created more clusters.", "Supports clustering for understandable traffic-pattern discovery."],
        ["Gupta, Kumar and Kumar (2025) [2]", "K-Means plus LR, DT, GB, RF and SVR for congestion analysis.", "Hybrid clustering improved several regression metrics; reported gains were dataset-specific.", "Motivates separating clusters from supervised comparisons."],
        ["Amin et al. (2025) [3]", "Eight supervised classifiers on balanced synthetic congestion scenarios.", "Hyperparameter optimisation changed comparative ranking.", "Shows why evaluation settings must be reported."],
        ["Arabiat, Hassan and Almomani (2024) [4]", "Hourly Amman data; LR, KNN, DT, RF, SVM and MLP classifiers.", "SVM was strongest in that local classification task.", "Demonstrates that best model depends on task and location."],
        ["Shaikh, Jaiswal and Sirsikar (2024) [5]", "Traffic-intensity regression comparison of tree and ensemble models.", "Ensembles outperformed single trees for their long-term task.", "Provides an ensemble baseline for the present test."],
        ["Sayed, Abdel-Hamid and Hefny (2023) [6]", "Review of ML and DL traffic-flow methods.", "Data quality, nonlinear patterns and computing limits shape model choice.", "Frames limitations and deployment requirements."],
    ], widths=[1.25, 1.75, 2.05, 1.85], font_size=7.6)
    heading(doc, "2.1 Comparative synthesis", 2)
    write_para(doc, "Three common lessons emerge. First, clustering can reveal meaningful operating states before a user looks at raw records [1, 2]. Second, an ensemble or deep model may improve accuracy in a particular task, but it is not automatically the best choice when interpretability and fast interaction are essential [3-6]. Third, reported accuracy, F1 score, RMSE and R-squared measure different targets. This report therefore compares four regression models on one held-out split and reports K-Means quality separately, instead of mixing incomparable metrics.")
    page_break(doc)

    # PAGE 4
    heading(doc, "3 Methodology")
    write_para(doc, "FlowSight follows a two-layer design. The offline layer generates a reproducible dataset, prepares features and fits the models. The online layer accepts a user scenario, validates it and requests one coordinated analysis from the Flask API. The React client then displays results as traffic groups, comparable records, contribution bars and rule paths.")
    add_image(doc, "fig3_system_flow.png", 6.65)
    caption(doc, "Figure 1. FlowSight interactive analytics flow.")
    heading(doc, "3.1 Dataset and features", 2)
    add_table(doc, ["Feature", "Type", "Meaning and project range"], [
        ["hour", "Integer", "Time of observation, 0 to 23."],
        ["traffic_volume", "Integer", "Vehicles observed, 0 to 3,200."],
        ["average_speed_kmph", "Numeric", "Observed average speed, 1 to 100 km/h."],
        ["rainfall_mm", "Numeric", "Rainfall amount, 0 to 30 mm."],
        ["incidents", "Integer", "Reported road incidents, 0 to 3."],
        ["road_lanes", "Integer", "Open lanes, 1 to 6."],
        ["delay_minutes", "Numeric target", "Recorded delay used for descriptive and evaluation analysis."],
    ], widths=[2.0, 1.25, 3.4], font_size=8.5)
    heading(doc, "3.2 Algorithms and roles", 2)
    write_para(doc, "K-Means is fitted to standardized volume, speed, rainfall, incidents and delay, with k = 3. The clusters are named by ascending mean delay: free flow, moderate congestion and severe congestion. KNN standardizes the six user-input features and retrieves the closest records; it is a retrieval method in the dashboard. A depth-limited Decision Tree learns an interpretable approximation of the K-Means cluster label. Linear Regression fits delay from the same six inputs and calculates coefficient times input difference against a user-selected reference. This last view explains associations; it is not displayed as a future-delay forecast.")
    page_break(doc)

    # PAGE 5
    heading(doc, "4 Data Cleaning and Preprocessing")
    write_para(doc, "Cleaning was implemented before model fitting and repeated at the API boundary. The generated data has a transparent schema, which makes validation reproducible. In a real sensor deployment, the same sequence would also include source reconciliation, time-zone alignment, sensor-health flags and treatment of missing periods.")
    add_table(doc, ["Step", "Action", "Reason"], [
        ["Schema check", "Confirm six inputs and one delay field with expected numeric type.", "Prevents missing or mislabeled variables entering a model."],
        ["Range validation", "Reject hour outside 0-23, lanes below 1, negative rain and non-finite values.", "Stops impossible values and model errors."],
        ["Integer validation", "Hour, volume, incidents and lanes must be whole numbers.", "Keeps operational inputs meaningful."],
        ["Duplicate and completeness review", "Generated seed produces 2,400 complete rows; no missing records are retained.", "Avoids silent row loss or accidental repetition."],
        ["Feature scaling", "StandardScaler is fitted on feature values for K-Means and KNN.", "Prevents vehicle volume from dominating distance calculations."],
        ["Train-test split", "75% training (1,800 rows) and 25% test (600 rows), random_state = 42.", "Keeps model comparison separate from fitting."],
    ], widths=[1.35, 3.25, 2.05], font_size=8.2)
    heading(doc, "4.1 Preprocessing rationale", 2)
    write_para(doc, "Scaling is critical for K-Means and KNN because both use distances. Without it, volume measured in hundreds or thousands would outweigh rainfall and incidents merely because of units. The regression and Decision Tree can accept raw values, but the comparison uses a standardized pipeline for Linear Regression and KNN so each model receives the same six explanatory variables. No target leakage is used: delay is the regression target and is not included in the six user inputs used to retrieve neighbours.")
    heading(doc, "4.2 Data quality summary", 2)
    write_para(doc, f"The final analytic sample contains {facts['kmeans']['observations']:,} rows. The reproducible split produced {facts['kmeans']['train_records']:,} training records and {facts['kmeans']['test_records']:,} held-out test records. The absence of missing values is a property of this simulated dataset, not evidence that live sensor streams are clean. The production recommendation is to log rejected API inputs and monitor missingness by road segment and time interval.")
    add_image(doc, "fig4_group_delay.png", 5.35)
    caption(doc, "Figure 2. Delay profile of the K-Means groups. Labels are descriptive names assigned after clustering by average delay.")
    page_break(doc)

    # PAGE 6
    heading(doc, "5 Results and Discussion")
    write_para(doc, "K-Means formed three groups with a silhouette score of 0.451 and a Davies-Bouldin index of 1.014. The silhouette value indicates moderate separation for this compact classroom dataset; it is useful for checking whether the three-group display is reasonable, but it is not an accuracy percentage. The discovered groups contained 1,450 free-flow, 787 moderate and 163 severe observations.")
    add_image(doc, "fig1_clusters.png", 6.35)
    caption(doc, "Figure 3. Speed-delay relationship coloured by K-Means group. The overlap is expected because traffic conditions vary continuously.")
    heading(doc, "5.1 Interpretation of the four application views", 2)
    add_table(doc, ["View", "What the user sees", "Correct interpretation"], [
        ["K-Means", "Three traffic-group cards and group chart.", "A descriptive segmentation of the simulated observations."],
        ["KNN", "The closest 10 to 100 records and their average delay/speed.", "Evidence from comparable stored observations, not a forecast."],
        ["Decision Tree", "A small set of readable thresholds and a group label.", "A simplified explanation of the K-Means partition."],
        ["Linear Regression", "Positive and negative contribution bars relative to a reference.", "An association under the fitted linear model while other inputs are held fixed."],
    ], widths=[1.25, 2.95, 2.45], font_size=8.2)
    heading(doc, "5.2 Discussion", 2)
    write_para(doc, "The dashboard deliberately links each visual to a specific algorithm. When rainfall or volume changes, the nearest-record set changes, so the cards, chart and table change together. The linear bars answer a different question: how the current input differs from an example under a fitted linear relationship. Keeping these meanings separate makes the demo easier to defend in a viva. A small change may leave the nearest set unchanged; the interface does not force the graph to move when the data does not support a change.")
    page_break(doc)

    # PAGE 7
    heading(doc, "6 State of the Art Models Comparison")
    write_para(doc, "The literature includes ensemble, support-vector and deep-learning approaches [2-6]. To keep a fair classroom evaluation, this report compares four supervised regression baselines on exactly the same 1,800/600 train-test split and six inputs. Random Forest is added only as a comparison baseline; the deployed FlowSight views remain Linear Regression, KNN, Decision Tree and K-Means.")
    add_table(doc, ["Model", "MAE min", "RMSE min", "R squared", "Train ms", "Test inference ms"], [
        [row["model"], row["MAE (min)"], row["RMSE (min)"], row["R2"], row["Train time (ms)"], row["Test inference (ms)"]]
        for row in metrics
    ], widths=[1.55, .82, .9, .82, 1.05, 1.2], font_size=8.0)
    caption(doc, "Table 1. Held-out regression performance. Lower MAE/RMSE is better; higher R squared is better. Timings depend on the local computer and are included only as relative implementation evidence.")
    add_image(doc, "fig2_model_comparison.png", 6.35)
    caption(doc, "Figure 4. Comparable performance metrics on the held-out test set.")
    heading(doc, "6.1 Findings", 2)
    write_para(doc, "Linear Regression produced the lowest RMSE (3.424 minutes) and the highest R-squared (0.9605) in this synthetic data. That result is plausible because the data generator contains approximately additive relationships. KNN and Random Forest were close in error, while the single Decision Tree was fastest at test inference but less accurate. This does not prove that linear models are best for real traffic: the cited studies report settings where SVM, tree ensembles or deep models perform better [2-6]. It shows why the evaluation must match the dataset, target and deployment purpose.")
    heading(doc, "6.2 Clustering quality", 2)
    write_para(doc, f"K-Means is not scored with RMSE because it has no target label. Its quality measures are silhouette = {facts['kmeans']['silhouette']} (higher indicates more compact and separated groups) and Davies-Bouldin = {facts['kmeans']['davies_bouldin']} (lower is preferred). These measures justify the three-group exploration view rather than claiming classifier accuracy.")
    page_break(doc)

    # PAGE 8
    heading(doc, "7 User Interface Development and Deployment")
    write_para(doc, "The interface is built with React and Vite, and the analytics service is built with Flask. The browser communicates with the local API through same-origin deployment; Flask serves the built React files and endpoints such as POST /api/analyze. Vite can also proxy /api during development. The design avoids a submit button: inputs are debounced and every valid change triggers a fresh coordinated analysis.")
    dashboard = ROOT / "dashboard-preview.png"
    if dashboard.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(dashboard), width=Inches(5.0))
        caption(doc, "Figure 5. FlowSight dashboard showing editable inputs, discovered groups, chart controls and a comparable-record summary.")
    heading(doc, "7.1 Interaction and accessibility", 2)
    add_table(doc, ["Requirement", "Implemented behaviour"], [
        ["Editable inputs", "Numeric fields and keyboard-accessible sliders for all six traffic variables."],
        ["Immediate feedback", "250 ms debounce, live status text, validation messages and retry state."],
        ["Interactive figures", "Chart metric toggle; group cards and bars filter the records table."],
        ["Explainability", "Clicking a factor opens plain-language linear explanation; the tree lists actual path rules."],
        ["Responsive use", "Desktop sidebar becomes a single-column input layout on smaller screens."],
    ], widths=[1.55, 5.0], font_size=8.4)
    heading(doc, "7.2 Deployment procedure", 2)
    write_para(doc, "For a local demonstration, train the model artifacts with python backend/train_models.py, build the client with npm run build inside frontend, then run python backend/app.py from the project root. The browser opens http://127.0.0.1:5000. For development, npm run dev starts Vite at port 5173 and proxies API calls to Flask. A production deployment should use authenticated HTTPS hosting, city-approved datasets, monitoring and a scheduled model-refresh policy.")
    page_break(doc)

    # PAGE 9
    heading(doc, "8 Conclusion and Future Work")
    write_para(doc, "FlowSight meets the project goal of making traffic analytics simple to demonstrate. It combines four requested algorithms in one interaction: K-Means creates understandable groups, KNN retrieves comparable records, a Decision Tree turns a group assignment into short rules, and Linear Regression explains change relative to a reference. The application is more defensible than a static dashboard because every displayed group count, average, record table and contribution responds to the selected input scenario.")
    heading(doc, "8.1 Main findings", 2)
    add_table(doc, ["Area", "Evidence from this project"], [
        ["Data preparation", "Six validated user-input features and 2,400 complete reproducible observations."],
        ["Analytics", "Three K-Means groups; silhouette 0.451; KNN preserves the actual closest observation."],
        ["Regression comparison", "Four models evaluated on a 600-row held-out set; Linear Regression RMSE 3.424 min."],
        ["UI and deployment", "React dashboard with interactive cards, chart filters, record reuse and Flask API."],
        ["Testing", "Backend tests cover changed scenarios, group aggregates, exact nearest record, reference logic and invalid inputs."],
    ], widths=[1.85, 4.7], font_size=8.4)
    heading(doc, "8.2 Limitations", 2)
    write_para(doc, "The central limitation is data realism. The data is simulated, so the reported performance measures demonstrate internal consistency rather than transportability to a specific road network. The linear hour feature also does not encode midnight wraparound or richer temporal patterns. K-Means labels are descriptive and can overlap. The Decision Tree is an explanation surrogate and can disagree with the dominant KNN group. Finally, model timing depends on the local hardware and is not a cloud scalability benchmark.")
    heading(doc, "8.3 Future work", 2)
    write_para(doc, "The next version should ingest authorised real sensor, weather and incident feeds; add time-lagged and road-network features; compare calibrated gradient boosting and spatiotemporal neural models; evaluate across locations and seasons; and provide dataset drift alerts. A human reviewer should approve any operational traffic-management decision. These steps would move the project from an educational analytics dashboard toward a reliable decision-support prototype.")
    heading(doc, "Viva summary", 2)
    write_para(doc, "In a presentation, start by changing rainfall or vehicle volume. Point out that KNN selects new similar records, K-Means re-counts their traffic groups, Linear Regression explains the difference from a reference, and the Decision Tree gives simple rules. Finish by stating that these are simulated records and the dashboard analyses patterns instead of promising a real-world forecast.")
    page_break(doc)

    # PAGE 10
    heading(doc, "9 References")
    references = [
        "[1] Lahijanian, Z., Isaenko, N., and Fusco, G. (2025). Empirical analysis of traffic patterns: Leveraging machine learning techniques for in-depth insights. Transportation Research Procedia, 90, 250-257. https://doi.org/10.1016/j.trpro.2025.06.065",
        "[2] Gupta, S., Kumar, A., and Kumar, A. (2025). Analysis of traffic flow congestion by integrating supervised machine learning with K-mean clustering. SN Computer Science, 6, 255. https://doi.org/10.1007/s42979-025-03761-4",
        "[3] Amin, G. A. H., Hassan, S. A., Muhammed, R. K., Saeed, M. S., and Ahmed, T. N. (2025). A comparative analysis of machine learning algorithms for urban traffic congestion prediction: The impact of hyperparameter optimization. Academic Journal of International University of Erbil. https://journal.iue.edu.krd/index.php/public/article/view/238",
        "[4] Arabiat, A., Hassan, M., and Almomani, O. (2024). WEKA-based machine learning for traffic congestion prediction in Amman City. IAES International Journal of Artificial Intelligence, 13(4), 4422-4434. https://doi.org/10.11591/ijai.v13.i4.pp4422-4434",
        "[5] Shaikh, A., Jaiswal, V., and Sirsikar, S. (2024). Decision-tree-based ensemble learning models for long-term traffic intensity forecasting and analysis of congestion treatment strategies. International Journal on Recent and Innovation Trends in Computing and Communication, 11(10), 2740-2744. https://www.ijritcc.org/index.php/ijritcc/article/view/10303",
        "[6] Sayed, S. A., Abdel-Hamid, Y., and Hefny, H. A. (2023). Artificial intelligence-based traffic flow prediction: A comprehensive review. Journal of Electrical Systems and Information Technology, 10, 13. https://doi.org/10.1186/s43067-023-00081-6",
        "[7] Pereira, P., Sampaio, J., and Fonseca, J. (2023). A survey on traffic flow prediction and classification. Intelligent Systems with Applications, 20, 200268. https://doi.org/10.1016/j.iswa.2023.200268",
        "[8] Pedregosa, F., Varoquaux, G., Gramfort, A., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "[9] React. (2026). React documentation. https://react.dev/",
        "[10] Pallets. (2026). Flask documentation. https://flask.palletsprojects.com/",
    ]
    for reference in references:
        p = write_para(doc, reference, size=9.3, after=7)
        p.paragraph_format.left_indent = Inches(.18)
        p.paragraph_format.first_line_indent = Inches(-.18)
    heading(doc, "Appendix A Reproducibility Note", 2)
    write_para(doc, "Dataset generation uses NumPy random seed 42. K-Means uses three clusters, n_init = 25 and random_state = 42. The Decision Tree has maximum depth 3 in the deployed explainer. The comparison split uses test_size = 0.25 and random_state = 42. The included test metrics can therefore be reproduced from the project files. Literature performance numbers are not pooled because studies use different cities, targets, splits and metrics.", size=9.5)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
